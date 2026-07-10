# -*- coding: utf-8 -*-
"""
relatorio_semanal.py
─────────────────────────────────────────────────────────────────────────
Geração automática do "Relatório Semanal" (.pptx) no modelo Grid Co.,
a partir das ocorrências já registradas na planilha do Painel de Falhas.

Não usa nenhuma API de IA — o texto é montado por regras determinísticas
em cima dos campos: Equipamento, Falha e Status (resumido, sem despejar
o histórico bruto de ações).
"""

import os
import re
import copy
import unicodedata
from io import BytesIO
from datetime import datetime

from pptx import Presentation
from pptx.oxml.ns import qn

# ── Configuração ───────────────────────────────────────────────────────────

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "templates", "modelo_relatorio_semanal.pptx")

# Índices das colunas da planilha principal (0-based, iguais ao app.py / CAMPO_COL)
COL_ID, COL_CLIENTE, COL_USINA, COL_EQUIP = 0, 1, 2, 3
COL_FALHA, COL_CAUSA, COL_IMPACTADOS, COL_ACAO = 4, 5, 6, 7
COL_STATUS, COL_TICKET, COL_OS, COL_HISTORICO = 8, 9, 10, 11
COL_DATA_ABERTURA = 12
COL_DATA_FECHAMENTO = 20  # coluna U (21ª, 1-based) = índice 20 (0-based)

STATUS_CONCLUIDO = ["concluído", "concluido", "resolvido", "fechado", "resolved", "closed"]

# Categorias "bonitas" para exibir no relatório (padrões regex, avaliados em ordem —
# o primeiro que bater define a categoria). Casa tanto nomes por extenso ("Inversor")
# quanto códigos curtos usados em campo ("Inv-18", "NVR", "Chave Seccionadora Skid 1").
PADROES_CATEGORIA = [
    (r"\binv[\s\-]*\d+|inversor", "INVERSORES"),
    (r"tracker|\btcu\b", "TRACKERS / TCU"),
    (r"fusive|fusíve", "FUSÍVEIS"),
    (r"transformador", "TRANSFORMADORES"),
    (r"switchgear|chave seccionadora|disjuntor|religador", "SWITCHGEAR"),
    (r"\bstring\b|modulo|módulo|otimizador", "STRINGS / MÓDULOS"),
    (r"usina desligad|usina offline|desligamento|ufv desligad|usina parad", "DESLIGAMENTOS"),
    (r"comunica|scada|cctv|\bnvr\b|camera|câmera|speed dome|igate", "COMUNICAÇÕES / SCADA / CCTV"),
    (r"sensor|piranometro|piranômetro|solarimetric|estacao solarimetrica|estação solarimétrica", "SENSORES"),
    (r"emergenc", "EMERGÊNCIAS"),
    (r"operac|opera[çc][aã]o", "OPERAÇÕES"),
    (r"termografia|termograf", "TERMOGRAFIA"),
    (r"restart|religamento", "RESTART / RELIGAMENTO"),
    (r"ronda", "RONDAS SEMANAIS"),
    (r"exaustor", "EXAUSTORES"),
    (r"vegeta", "CONTROLE DE VEGETAÇÃO"),
]


# ── Utilidades de texto/data ────────────────────────────────────────────────

def _norm(s):
    s = unicodedata.normalize("NFKD", (s or "").lower())
    s = s.encode("ascii", "ignore").decode("ascii")
    return s.strip()


def _rotulo_categoria(equip_bruto):
    n = _norm(equip_bruto)
    for padrao, rotulo in PADROES_CATEGORIA:
        if re.search(padrao, n):
            return rotulo
    return (equip_bruto or "OUTRAS OCORRÊNCIAS").strip().upper()


def _parse_data(txt):
    """Aceita 'dd/mm/aaaa[ hh:mm[:ss]]' -> datetime, ou None."""
    if not txt:
        return None
    txt = txt.strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(txt, fmt)
        except ValueError:
            continue
    return None


def _is_concluido(status):
    s = (status or "").lower()
    return any(x in s for x in STATUS_CONCLUIDO)


def _ticket_valido(ticket):
    """Só considera 'chamado real' se houver um número/identificador de verdade."""
    t = (ticket or "").strip().lower()
    return t not in ("", "n/a", "na", "-", "--", "nao", "não", "sem chamado", "sem ticket", "s/n")


# ── Coleta e agrupamento dos dados ──────────────────────────────────────────

def coletar_ocorrencias_semana(todos_valores, cliente, data_inicio, data_fim):
    """
    todos_valores: retorno de ws.get_all_values() (linha 0 = cabeçalho)
    cliente: nome do cliente (comparação por 'contains', case-insensitive)
    data_inicio / data_fim: datetime (00:00 do dia inicial até 23:59 do dia final)

    Retorna dict: {"categoria": {"concluidas": [linha,...], "abertas": [linha,...]}, ...}
    """
    cliente_norm = _norm(cliente)
    grupos = {}

    for row in todos_valores[1:]:
        if len(row) <= COL_DATA_FECHAMENTO:
            row = row + [""] * (COL_DATA_FECHAMENTO + 1 - len(row))

        if cliente_norm not in _norm(row[COL_CLIENTE]):
            continue

        dt_abertura = _parse_data(row[COL_DATA_ABERTURA])
        dt_fecham = _parse_data(row[COL_DATA_FECHAMENTO])
        concluida = _is_concluido(row[COL_STATUS])

        fechou_na_semana = dt_fecham and data_inicio <= dt_fecham <= data_fim
        abriu_na_semana = dt_abertura and data_inicio <= dt_abertura <= data_fim
        em_aberto_backlog = not concluida

        if not (fechou_na_semana or abriu_na_semana or em_aberto_backlog):
            continue

        categoria = _rotulo_categoria(row[COL_EQUIP])
        grupos.setdefault(categoria, {"concluidas": [], "abertas": []})

        if fechou_na_semana:
            grupos[categoria]["concluidas"].append(row)
        elif em_aberto_backlog:
            grupos[categoria]["abertas"].append(row)

    return {k: v for k, v in grupos.items() if v["concluidas"] or v["abertas"]}


def listar_usinas_cliente(todos_valores, cliente):
    """Lista todas as usinas do cliente (histórico completo, não só a semana)."""
    cliente_norm = _norm(cliente)
    usinas = []
    vistos = set()
    for row in todos_valores[1:]:
        if len(row) <= COL_USINA:
            continue
        if cliente_norm not in _norm(row[COL_CLIENTE]):
            continue
        usina = row[COL_USINA].strip()
        chave = _norm(usina)
        if usina and chave not in vistos:
            vistos.add(chave)
            usinas.append(usina)
    return sorted(usinas)


def coletar_chamados_abertos(todos_valores, cliente):
    """Ocorrências do cliente com número de chamado/ticket fabricante válido e ainda não concluídas."""
    cliente_norm = _norm(cliente)
    chamados = []
    for row in todos_valores[1:]:
        if len(row) <= COL_DATA_FECHAMENTO:
            row = row + [""] * (COL_DATA_FECHAMENTO + 1 - len(row))
        if cliente_norm not in _norm(row[COL_CLIENTE]):
            continue
        ticket = row[COL_TICKET].strip()
        if not _ticket_valido(ticket) or _is_concluido(row[COL_STATUS]):
            continue
        chamados.append(row)
    return chamados


# ── Geração de conteúdo (offline, sem IA) — parágrafos com negrito por usina ─

def _linha_ocorrencia(row):
    """Linha resumida: Equipamento – Falha. Status: X. (sem despejar causa/ação/histórico)"""
    equip = row[COL_EQUIP].strip() or "Equipamento não informado"
    falha = row[COL_FALHA].strip() or "Falha não especificada"
    if len(falha) > 180:
        falha = falha[:177].rstrip() + "..."
    status = row[COL_STATUS].strip() or "Em aberto"
    return f"{equip} – {falha}. Status: {status}."


def gerar_paragrafos_categoria(categoria, dados):
    """
    Agrupa as ocorrências de uma categoria por usina (nome da usina em negrito,
    seguido de um bullet resumido por ocorrência).
    Retorna lista de dicts: {"texto": str, "bold": bool}
    """
    todas = list(dados["concluidas"]) + list(dados["abertas"])
    por_usina = {}
    for row in todas:
        usina = row[COL_USINA].strip() or "Usina não informada"
        por_usina.setdefault(usina, []).append(row)

    paragrafos = []
    for usina in sorted(por_usina.keys()):
        paragrafos.append({"texto": usina, "bold": True})
        for row in por_usina[usina]:
            paragrafos.append({"texto": "• " + _linha_ocorrencia(row), "bold": False})
    return paragrafos


def gerar_paragrafos_chamados(chamados):
    """Chamados com ticket válido, agrupados por usina em negrito."""
    if not chamados:
        return [{"texto": "Nenhum chamado em aberto nesta semana.", "bold": False}]

    por_usina = {}
    for row in chamados:
        usina = row[COL_USINA].strip() or "Usina não informada"
        por_usina.setdefault(usina, []).append(row)

    paragrafos = []
    for usina in sorted(por_usina.keys()):
        paragrafos.append({"texto": usina, "bold": True})
        for row in por_usina[usina]:
            ticket = row[COL_TICKET].strip()
            resumo = row[COL_FALHA].strip() or row[COL_CAUSA].strip() or "Sem descrição detalhada"
            status = row[COL_STATUS].strip() or "Em análise"
            paragrafos.append({"texto": f"• {ticket} -> {resumo} -> {status}.", "bold": False})
    return paragrafos


# Layout real da aba "Zeladoria" (gid 987654321): linha 1 = título dos grupos,
# linha 2 = subcabeçalhos, dados a partir da linha 3.
# Colunas (0-based): A Cliente, B Usina, depois 4 grupos de 4 colunas cada
# (Última Data, Próxima Data, Quantidade, Status): Roçada, Poda Química, Lavagem dos Módulos, Controle de Pragas.
ZEL_COL_CLIENTE, ZEL_COL_USINA = 0, 1
ZEL_GRUPOS = [
    ("Roçada",               2),
    ("Poda Química",         6),
    ("Lavagem dos Módulos", 10),
    ("Controle de Pragas",  14),
]


def coletar_zeladoria(zeladoria_valores, cliente):
    """
    zeladoria_valores: ws.get_all_values() da aba Zeladoria (linhas 1-2 = cabeçalho).
    Retorna lista de dicts: [{"usina": ..., "grupos": [{"nome":, "status":, "ultima_data":}, ...]}, ...]
    """
    cliente_norm = _norm(cliente)
    resultado = []
    for row in zeladoria_valores[2:]:
        if len(row) <= ZEL_GRUPOS[-1][1] + 3:
            row = row + [""] * (ZEL_GRUPOS[-1][1] + 4 - len(row))
        if not row[ZEL_COL_USINA].strip():
            continue
        if cliente_norm not in _norm(row[ZEL_COL_CLIENTE]):
            continue

        grupos_usina = []
        for nome, col_ini in ZEL_GRUPOS:
            ultima_data = row[col_ini].strip()
            status = row[col_ini + 3].strip()
            grupos_usina.append({"nome": nome, "status": status, "ultima_data": ultima_data})

        resultado.append({"usina": row[ZEL_COL_USINA].strip(), "grupos": grupos_usina})
    return resultado


def gerar_paragrafos_zeladoria(zeladoria_usinas):
    """Uma usina em negrito, seguida de um bullet por frente (Roçada, Poda, Lavagem, Pragas)."""
    if not zeladoria_usinas:
        return [{"texto": "Nenhuma usina encontrada na aba de Zeladoria para este cliente.", "bold": False}]

    paragrafos = []
    for item in zeladoria_usinas:
        paragrafos.append({"texto": item["usina"], "bold": True})
        for g in item["grupos"]:
            status = g["status"] or "Sem informação"
            if g["ultima_data"] and status.lower() not in ("sem informação", "sem info"):
                linha = f"• {g['nome']}: {status} (última {g['ultima_data']})"
            else:
                linha = f"• {g['nome']}: {status}"
            paragrafos.append({"texto": linha, "bold": False})
    return paragrafos


# ── Clonagem de slides (mantém 100% do estilo do modelo Grid Co.) ──────────

def _duplicate_slide(prs, index):
    """Duplica um slide existente do modelo (mesmo layout, cores, fontes, formas)."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)

    id_map = {}
    for rId, rel in source.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        new_rid = dest.part.rels._add_relationship(rel.reltype, rel._target, rel.is_external)
        id_map[rId] = new_rid

    for shp in source.shapes:
        newel = copy.deepcopy(shp._element)
        for el in newel.iter():
            for attr in ("embed", "link", "id"):
                full = qn(f"r:{attr}")
                val = el.get(full)
                if val and val in id_map:
                    el.set(full, id_map[val])
        dest.shapes._spTree.append(newel)
    return dest


def _extrair_formato_base(shape):
    """Extrai fonte/tamanho/itálico/cor do 1º run existente na shape (pra reaplicar depois)."""
    tf = shape.text_frame
    primeiro_par = tf.paragraphs[0]
    if not primeiro_par.runs:
        return None
    font = primeiro_par.runs[0].font
    cor_rgb = None
    try:
        if font.color and font.color.type is not None:
            cor_rgb = font.color.rgb  # levanta AttributeError se for cor de tema (schemeClr)
    except AttributeError:
        cor_rgb = None
    return (font.name, font.size, font.italic, cor_rgb)


def _set_text_preservando_estilo(shape, novo_texto):
    """Substitui o texto de um text box mantendo a formatação (incl. negrito) do 1º run."""
    tf = shape.text_frame
    primeiro_par = tf.paragraphs[0]
    bold_original = primeiro_par.runs[0].font.bold if primeiro_par.runs else None
    fmt = _extrair_formato_base(shape)

    tf.clear()
    linhas = novo_texto.split("\n")
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = linha
        if fmt:
            if fmt[0]:
                run.font.name = fmt[0]
            if fmt[1]:
                run.font.size = fmt[1]
            if fmt[2] is not None:
                run.font.italic = fmt[2]
            if fmt[3]:
                run.font.color.rgb = fmt[3]
        if bold_original is not None:
            run.font.bold = bold_original


def _set_paragrafos_com_estilo(shape, paragrafos):
    """
    Substitui o conteúdo de um text box por uma lista de parágrafos, permitindo
    alternar negrito por parágrafo (usado para destacar nome de usina/categoria).
    paragrafos: lista de dicts {"texto": str, "bold": bool}
    """
    fmt = _extrair_formato_base(shape)
    tf = shape.text_frame
    tf.clear()

    if not paragrafos:
        paragrafos = [{"texto": "", "bold": False}]

    for i, p in enumerate(paragrafos):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = par.add_run()
        run.text = p["texto"]
        if fmt:
            if fmt[0]:
                run.font.name = fmt[0]
            if fmt[1]:
                run.font.size = fmt[1]
            if fmt[2] is not None:
                run.font.italic = fmt[2]
            if fmt[3]:
                run.font.color.rgb = fmt[3]
        run.font.bold = bool(p.get("bold", False))


def _remover_fotos(slide):
    """Remove os placeholders de foto (o gerador não tem fotos de campo)."""
    for shp in list(slide.shapes):
        if shp.shape_type == 13 and (shp.name or "").lower().startswith("imagem"):
            shp._element.getparent().remove(shp._element)


def _find_shape(slide, contem_texto):
    for shp in slide.shapes:
        if shp.has_text_frame and contem_texto.lower() in shp.text_frame.text.lower():
            return shp
    return None


# ── Montagem do PPTX final ──────────────────────────────────────────────────

def _tamanho_paragrafos(paragrafos):
    return sum(len(p["texto"]) for p in paragrafos)


def agrupar_em_paginas(grupos, orcamento_chars=1400):
    """
    Agrupa categorias em páginas por orçamento de caracteres. Retorna lista de
    páginas, cada página = lista de tuplas (categoria, paragrafos).
    """
    blocos = [(cat, gerar_paragrafos_categoria(cat, dados)) for cat, dados in grupos.items()]
    paginas, atual, atual_len = [], [], 0
    for cat, paragrafos in blocos:
        tam = _tamanho_paragrafos(paragrafos) + len(cat) + 10
        if atual and (atual_len + tam > orcamento_chars):
            paginas.append(atual)
            atual, atual_len = [], 0
        atual.append((cat, paragrafos))
        atual_len += tam
    if atual:
        paginas.append(atual)
    return paginas


def gerar_relatorio_pptx(cliente, semana_num, data_label, grupos, chamados=None, zeladoria_usinas=None):
    """
    grupos: dict categoria -> {"concluidas": [...], "abertas": [...]}
    semana_num: número da semana do ano (baseado na data final do período)
    data_label: data final formatada (dd/mm/aaaa)
    chamados: lista de linhas com chamado/ticket válido em aberto (opcional)
    zeladoria_usinas: retorno de coletar_zeladoria() (opcional)
    Retorna BytesIO() pronto para download.
    """
    prs = Presentation(TEMPLATE_PATH)

    # --- Slide 1: Capa — só cliente, semana e data -------------------------
    capa = _duplicate_slide(prs, 0)
    shp_titulo = _find_shape(capa, "O&M")
    if shp_titulo:
        _set_text_preservando_estilo(
            shp_titulo, f"O&M – {cliente}\nSemana {semana_num}\n{data_label}"
        )

    # --- Slide 2: Ata da reunião (tópicos/categorias da semana + extras) --
    topicos = list(grupos.keys())
    if chamados:
        topicos.append("Chamados em aberto")
    if zeladoria_usinas:
        topicos.append("Zeladoria")
    ata = _duplicate_slide(prs, 1)
    shp_lista = _find_shape(ata, "Desligamentos")
    if shp_lista:
        _set_text_preservando_estilo(shp_lista, "\n".join(topicos))

    # --- Uma página por grupo de categorias de ocorrências (empacotadas) ---
    for pagina in agrupar_em_paginas(grupos):
        novo = _duplicate_slide(prs, 2)
        _remover_fotos(novo)

        shp_categoria = _find_shape(novo, "DESLIGAMENTOS")
        shp_corpo = _find_shape(novo, "Foram registradas")

        if len(pagina) == 1:
            categoria, paragrafos = pagina[0]
            if shp_categoria:
                _set_text_preservando_estilo(shp_categoria, categoria)
            if shp_corpo:
                _set_paragrafos_com_estilo(shp_corpo, paragrafos)
        else:
            if shp_categoria:
                _set_text_preservando_estilo(shp_categoria, "OCORRÊNCIAS DA SEMANA")
            if shp_corpo:
                combinados = []
                for cat, paragrafos in pagina:
                    combinados.append({"texto": cat, "bold": True})
                    combinados.extend(paragrafos)
                _set_paragrafos_com_estilo(shp_corpo, combinados)

    # --- Página: Chamados em aberto (só com ticket/case válido) ------------
    if chamados:
        pg_chamados = _duplicate_slide(prs, 2)
        _remover_fotos(pg_chamados)
        shp_cat = _find_shape(pg_chamados, "DESLIGAMENTOS")
        shp_corpo = _find_shape(pg_chamados, "Foram registradas")
        if shp_cat:
            _set_text_preservando_estilo(shp_cat, "CHAMADOS EM ABERTO")
        if shp_corpo:
            _set_paragrafos_com_estilo(shp_corpo, gerar_paragrafos_chamados(chamados))

    # --- Página: Zeladoria --------------------------------------------------
    if zeladoria_usinas:
        pg_zeladoria = _duplicate_slide(prs, 2)
        _remover_fotos(pg_zeladoria)
        shp_cat = _find_shape(pg_zeladoria, "DESLIGAMENTOS")
        shp_corpo = _find_shape(pg_zeladoria, "Foram registradas")
        if shp_cat:
            _set_text_preservando_estilo(shp_cat, "ZELADORIA")
        if shp_corpo:
            _set_paragrafos_com_estilo(shp_corpo, gerar_paragrafos_zeladoria(zeladoria_usinas))

    # --- Reordena o deck: capa nova -> ata nova -> conteúdo -> contato -----
    xml_slides = prs.slides._sldIdLst
    todos_els = list(xml_slides)
    contato_el = todos_els[6]
    capa_el = todos_els[7]
    ata_el = todos_els[8]
    conteudo_els = todos_els[9:]

    for e in todos_els:
        xml_slides.remove(e)
    for e in [capa_el, ata_el] + conteudo_els + [contato_el]:
        xml_slides.append(e)

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
